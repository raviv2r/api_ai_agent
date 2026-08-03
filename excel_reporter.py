import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from typing import Dict, Any, List


class ExcelReportGenerator:
    """Generate comprehensive Excel reports for API testing"""

    def __init__(self, test_results: List[Dict[str, Any]], test_cases: List[Dict[str, Any]],
                 execution_details: Dict[str, Any], custom_script_results: List[Dict[str, Any]] = None):
        self.test_results = test_results
        self.test_cases = test_cases
        self.execution_details = execution_details
        self.custom_script_results = custom_script_results or []
        self.workbook = openpyxl.Workbook()
        self.workbook.remove(self.workbook.active)  # Remove default sheet

    def generate_report(self) -> str:
        """Generate complete Excel report with multiple sheets"""
        # Create sheets
        self._create_summary_sheet()
        self._create_test_results_sheet()
        self._create_test_cases_sheet()
        self._create_execution_details_sheet()
        self._create_custom_scripts_sheet()

        # Save report
        filename = f"API_Test_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        self.workbook.save(filename)
        return filename

    def _create_summary_sheet(self):
        """Create summary sheet with statistics"""
        ws = self.workbook.create_sheet("Summary", 0)

        # Calculate statistics
        total_tests = len(self.test_results)
        passed_tests = sum(1 for test in self.test_results if test.get("passed", False))
        failed_tests = total_tests - passed_tests
        pass_percentage = (passed_tests / total_tests * 100) if total_tests > 0 else 0
        fail_percentage = (failed_tests / total_tests * 100) if total_tests > 0 else 0

        # Count by category
        by_category = {}
        for test in self.test_results:
            category = test.get("category", "Other")
            if category not in by_category:
                by_category[category] = {"passed": 0, "failed": 0}
            if test.get("passed"):
                by_category[category]["passed"] += 1
            else:
                by_category[category]["failed"] += 1

        # Title
        ws["A1"] = "API TEST EXECUTION SUMMARY"
        ws["A1"].font = Font(bold=True, size=16)
        ws.merge_cells("A1:C1")

        # Execution metadata
        row = 3
        ws[f"A{row}"] = "Execution Details"
        ws[f"A{row}"].font = Font(bold=True, size=12)

        row += 1
        ws[f"A{row}"] = "API Endpoint:"
        ws[f"B{row}"] = self.execution_details.get("url", "N/A")
        row += 1
        ws[f"A{row}"] = "Request Method:"
        ws[f"B{row}"] = self.execution_details.get("method", "N/A")
        row += 1
        ws[f"A{row}"] = "HTTP Status Code:"
        ws[f"B{row}"] = self.execution_details.get("status_code", "N/A")
        row += 1
        ws[f"A{row}"] = "Response Time:"
        ws[f"B{row}"] = f"{self.execution_details.get('response_time', 0)}s"
        row += 1
        ws[f"A{row}"] = "Execution Date & Time:"
        ws[f"B{row}"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        # Summary statistics
        row += 2
        ws[f"A{row}"] = "TEST SUMMARY"
        ws[f"A{row}"].font = Font(bold=True, size=12)

        row += 1
        # Headers
        headers = ["Metric", "Count", "Percentage"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col)
            cell.value = header
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Data rows
        row += 1
        metrics = [
            ("Total Tests", total_tests, "100%"),
            ("Passed Tests", passed_tests, f"{pass_percentage:.1f}%"),
            ("Failed Tests", failed_tests, f"{fail_percentage:.1f}%"),
        ]

        for metric, count, percentage in metrics:
            ws.cell(row=row, column=1).value = metric
            ws.cell(row=row, column=2).value = count
            ws.cell(row=row, column=3).value = percentage

            # Color coding
            if "Passed" in metric:
                ws.cell(row=row, column=2).fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            elif "Failed" in metric:
                ws.cell(row=row, column=2).fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

            row += 1

        # Tests by category
        row += 1
        ws[f"A{row}"] = "TESTS BY CATEGORY"
        ws[f"A{row}"].font = Font(bold=True, size=12)

        row += 1
        headers = ["Category", "Passed", "Failed", "Total"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col)
            cell.value = header
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        row += 1
        for category, stats in sorted(by_category.items()):
            ws.cell(row=row, column=1).value = category
            ws.cell(row=row, column=2).value = stats["passed"]
            ws.cell(row=row, column=3).value = stats["failed"]
            ws.cell(row=row, column=4).value = stats["passed"] + stats["failed"]

            # Color code passed/failed
            ws.cell(row=row, column=2).fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            if stats["failed"] > 0:
                ws.cell(row=row, column=3).fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

            row += 1

        # Adjust column widths
        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 15

    def _create_test_results_sheet(self):
        """Create detailed test results sheet"""
        ws = self.workbook.create_sheet("Test Results", 1)

        # Headers
        headers = [
            "Test ID",
            "Test Name",
            "Category",
            "Status",
            "Expected",
            "Actual",
            "Failure Reason",
            "HTTP Status",
            "Response Time (s)",
            "API Endpoint",
            "Method",
            "Execution Time"
        ]

        # Add headers with formatting
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col)
            cell.value = header
            cell.font = Font(bold=True, color="FFFFFF", size=11)
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Add test results
        for idx, test in enumerate(self.test_results, 2):
            ws.cell(row=idx, column=1).value = idx - 1
            ws.cell(row=idx, column=2).value = test.get("name", "N/A")
            ws.cell(row=idx, column=3).value = test.get("category", "Other")
            
            status = "✅ PASS" if test.get("passed", False) else "❌ FAIL"
            ws.cell(row=idx, column=4).value = status
            
            # Color code status
            if test.get("passed"):
                ws.cell(row=idx, column=4).fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            else:
                ws.cell(row=idx, column=4).fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

            ws.cell(row=idx, column=5).value = test.get("expected", "N/A")
            ws.cell(row=idx, column=6).value = test.get("actual", "N/A")
            ws.cell(row=idx, column=7).value = test.get("message", "N/A") if not test.get("passed") else ""
            ws.cell(row=idx, column=8).value = self.execution_details.get("status_code", "N/A")
            ws.cell(row=idx, column=9).value = self.execution_details.get("response_time", "N/A")
            ws.cell(row=idx, column=10).value = self.execution_details.get("url", "N/A")
            ws.cell(row=idx, column=11).value = self.execution_details.get("method", "N/A")
            ws.cell(row=idx, column=12).value = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

            # Set alignment and wrap text
            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=idx, column=col)
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        # Adjust column widths
        widths = [10, 25, 15, 12, 20, 20, 20, 12, 15, 30, 10, 20]
        for col, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = width

        # Freeze top row
        ws.freeze_panes = "A2"

    def _create_test_cases_sheet(self):
        """Create test cases sheet"""
        ws = self.workbook.create_sheet("Test Cases", 2)

        # Headers
        headers = [
            "Test Case ID",
            "Test Case Name",
            "Description",
            "Steps",
            "Expected Result"
        ]

        # Add headers
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col)
            cell.value = header
            cell.font = Font(bold=True, color="FFFFFF", size=11)
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Add test cases
        for idx, tc in enumerate(self.test_cases, 2):
            ws.cell(row=idx, column=1).value = tc.get("id", "N/A")
            ws.cell(row=idx, column=2).value = tc.get("name", "N/A")
            ws.cell(row=idx, column=3).value = tc.get("description", "N/A")
            
            # Format steps as numbered list
            steps = tc.get("steps", [])
            steps_text = "\n".join([f"{i}. {step}" for i, step in enumerate(steps, 1)])
            ws.cell(row=idx, column=4).value = steps_text
            
            ws.cell(row=idx, column=5).value = tc.get("expected_result", "N/A")

            # Set alignment and wrap text
            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=idx, column=col)
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        # Adjust column widths
        ws.column_dimensions["A"].width = 15
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 35
        ws.column_dimensions["D"].width = 40
        ws.column_dimensions["E"].width = 35

        # Set row height for better readability
        ws.row_dimensions[1].height = 30

        # Freeze top row
        ws.freeze_panes = "A2"

    def _create_execution_details_sheet(self):
        """Create execution details sheet"""
        ws = self.workbook.create_sheet("Execution Details", 3)

        ws["A1"] = "API TEST EXECUTION DETAILS"
        ws["A1"].font = Font(bold=True, size=14)
        ws.merge_cells("A1:B1")

        row = 3
        details = [
            ("API Endpoint", self.execution_details.get("url", "N/A")),
            ("Request Method", self.execution_details.get("method", "N/A")),
            ("HTTP Status Code", self.execution_details.get("status_code", "N/A")),
            ("Response Time (seconds)", self.execution_details.get("response_time", "N/A")),
            ("Execution Date & Time", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("Total Test Cases", len(self.test_cases)),
            ("Total Validations Executed", len(self.test_results)),
            ("Passed Validations", sum(1 for test in self.test_results if test.get("passed", False))),
            ("Failed Validations", sum(1 for test in self.test_results if not test.get("passed", False))),
        ]

        for label, value in details:
            ws[f"A{row}"] = label
            ws[f"A{row}"].font = Font(bold=True)
            ws[f"B{row}"] = value
            row += 1

        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 40

    def _create_custom_scripts_sheet(self):
        """Create a sheet for Post-response custom script results"""
        ws = self.workbook.create_sheet("Custom Script Results", 4)

        ws["A1"] = "POST-RESPONSE SCRIPT RESULTS"
        ws["A1"].font = Font(bold=True, size=14)
        ws.merge_cells("A1:F1")

        total   = len(self.custom_script_results)
        passed  = sum(1 for r in self.custom_script_results if r.get("passed") and not r.get("skipped"))
        failed  = sum(1 for r in self.custom_script_results if not r.get("passed") and not r.get("skipped"))
        skipped = sum(1 for r in self.custom_script_results if r.get("skipped"))

        row = 3
        for label, val in [("Total Scripts", total), ("Passed", passed),
                            ("Failed", failed), ("Skipped", skipped)]:
            ws.cell(row=row, column=1).value = label
            ws.cell(row=row, column=1).font = Font(bold=True)
            ws.cell(row=row, column=2).value = val
            row += 1

        row += 1
        col_headers = ["#", "Script Name", "Status", "Expected", "Actual", "Failure Reason"]
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        for col, h in enumerate(col_headers, 1):
            cell = ws.cell(row=row, column=col)
            cell.value = h
            cell.font = Font(bold=True, color="FFFFFF", size=11)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        row += 1
        for idx, r in enumerate(self.custom_script_results, 1):
            if r.get("skipped"):
                status_text = "SKIPPED"
                fill = PatternFill(start_color="FFFACD", end_color="FFFACD", fill_type="solid")
            elif r.get("passed"):
                status_text = "PASS"
                fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            else:
                status_text = "FAIL"
                fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

            ws.cell(row=row, column=1).value = idx
            ws.cell(row=row, column=2).value = r.get("name", "N/A")
            ws.cell(row=row, column=3).value = status_text
            ws.cell(row=row, column=3).fill = fill
            ws.cell(row=row, column=4).value = r.get("expected", "")
            ws.cell(row=row, column=5).value = r.get("actual", "")
            ws.cell(row=row, column=6).value = r.get("failure_reason", "")

            for col in range(1, 7):
                ws.cell(row=row, column=col).alignment = Alignment(wrap_text=True, vertical="top")
            row += 1

        for col, width in zip(range(1, 7), [6, 35, 12, 30, 30, 40]):
            ws.column_dimensions[get_column_letter(col)].width = width

        ws.freeze_panes = "A8"


class ReportDownloader:
    """Helper class for report download management"""

    @staticmethod
    def get_download_link(filename: str) -> str:
        """Generate download link for Excel report"""
        return f"Download Report: {filename}"
